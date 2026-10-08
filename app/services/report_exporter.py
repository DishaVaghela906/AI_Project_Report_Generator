"""
Report Export Service.

Generates final PDF or DOCX reports from stored chapter content.
Strict Academic Formatting:
  - Font: Times New Roman
  - Main Headings (Chapter Title / H1): 16 pt (Bold)
  - Sub Headings (H2 / H3): 14 pt (Bold)
  - Content (Body Text, Bullets, Table Cells): 12 pt
  - Content Alignment: Justified
  - Full Markdown Table & Code Block Support
"""
from __future__ import annotations

import re
import unicodedata
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from loguru import logger
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.core.config import settings

EXPORT_DIR = Path(settings.UPLOAD_DIR).parent / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)


# ─── Public API ───────────────────────────────────────────────────────────────

def export_to_pdf(
    project: dict,
    chapters: list[dict],
    include_chapters: list[int] | None = None,
    cover_page: bool = True,
) -> Path:
    """Generate a PDF report adhering strictly to Times New Roman academic styling."""
    filename = _safe_filename(project.get("title", "report")) + ".pdf"
    out_path = EXPORT_DIR / filename

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        rightMargin=54,  # ~0.75 inch
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = _build_pdf_styles()
    story = []

    if cover_page:
        story += _pdf_cover_page(project, styles)

    filtered = _filter_chapters(chapters, include_chapters)
    for chap in filtered:
        story += _pdf_chapter(chap, styles)

    doc.build(story)
    logger.success(f"PDF exported successfully: {out_path}")
    return out_path


def export_to_docx(
    project: dict,
    chapters: list[dict],
    include_chapters: list[int] | None = None,
    cover_page: bool = True,
) -> Path:
    """Generate a DOCX report adhering strictly to Times New Roman academic styling."""
    filename = _safe_filename(project.get("title", "report")) + ".docx"
    out_path = EXPORT_DIR / filename

    doc = Document()
    _set_docx_styles(doc)

    if cover_page:
        _docx_cover_page(doc, project)

    filtered = _filter_chapters(chapters, include_chapters)
    for chap in filtered:
        _docx_chapter(doc, chap)

    doc.save(str(out_path))
    logger.success(f"DOCX exported successfully: {out_path}")
    return out_path


# ─── PDF Styles & Builders ─────────────────────────────────────────────────────

def _build_pdf_styles():
    styles = getSampleStyleSheet()

    # Cover Page Styles
    styles.add(ParagraphStyle(
        name="CoverTitle", fontSize=20, leading=26, alignment=TA_CENTER,
        spaceAfter=14, fontName="Times-Bold", textColor=colors.HexColor("#1F1C18")
    ))
    styles.add(ParagraphStyle(
        name="CoverSubtitle", fontSize=14, leading=20, alignment=TA_CENTER,
        spaceAfter=8, fontName="Times-Roman", textColor=colors.HexColor("#333333")
    ))
    styles.add(ParagraphStyle(
        name="CoverMeta", fontSize=12, leading=18, alignment=TA_CENTER,
        spaceAfter=6, fontName="Times-Roman", textColor=colors.HexColor("#444444")
    ))

    # Main Heading: 16 pt (Bold)
    styles.add(ParagraphStyle(
        name="MainHeading", fontSize=16, leading=22, spaceAfter=12,
        spaceBefore=18, fontName="Times-Bold", alignment=TA_LEFT,
        textColor=colors.HexColor("#111111"), keepWithNext=True
    ))

    # Sub Headings: 14 pt (Bold)
    styles.add(ParagraphStyle(
        name="SubHeading", fontSize=14, leading=19, spaceAfter=8,
        spaceBefore=14, fontName="Times-Bold", alignment=TA_LEFT,
        textColor=colors.HexColor("#222222"), keepWithNext=True
    ))

    # Content Paragraph: 12 pt (Justified)
    styles.add(ParagraphStyle(
        name="BodyJustified", fontSize=12, leading=18, spaceAfter=10,
        fontName="Times-Roman", alignment=TA_JUSTIFY,
        textColor=colors.HexColor("#1F1C18")
    ))

    # Bullet Items: 12 pt (Justified)
    styles.add(ParagraphStyle(
        name="BulletJustified", fontSize=12, leading=18, spaceAfter=6,
        fontName="Times-Roman", alignment=TA_JUSTIFY, leftIndent=18,
        textColor=colors.HexColor("#1F1C18")
    ))

    # Table Cell Styles: 12 pt
    styles.add(ParagraphStyle(
        name="TableCell", fontSize=12, leading=16, fontName="Times-Roman",
        alignment=TA_LEFT, textColor=colors.HexColor("#1F1C18")
    ))
    styles.add(ParagraphStyle(
        name="TableCellBold", fontSize=12, leading=16, fontName="Times-Bold",
        alignment=TA_LEFT, textColor=colors.HexColor("#1F1C18")
    ))

    # Code Block Style: 11 pt Courier
    styles.add(ParagraphStyle(
        name="CodeBlock", fontSize=11, leading=15, spaceAfter=10,
        fontName="Courier", alignment=TA_LEFT, leftIndent=15,
        textColor=colors.HexColor("#2C2621"), backColor=colors.HexColor("#F4EFE6")
    ))

    return styles


def _pdf_cover_page(project: dict, styles) -> list:
    team_names = "; ".join(
        m.get("name", "") for m in (project.get("team_members") or [])
    )
    elements = [
        Spacer(1, 1.2 * inch),
        Paragraph(project.get("institution", "University Name"), styles["CoverSubtitle"]),
        Paragraph(project.get("department", "Department of Computer Science"), styles["CoverSubtitle"]),
        Spacer(1, 0.4 * inch),
        HRFlowable(width="85%", thickness=1.5, color=colors.HexColor("#1F1C18")),
        Spacer(1, 0.4 * inch),
        Paragraph("PROJECT REPORT", styles["CoverSubtitle"]),
        Spacer(1, 0.2 * inch),
        Paragraph(project.get("title", "Project Title"), styles["CoverTitle"]),
        Spacer(1, 0.4 * inch),
        HRFlowable(width="85%", thickness=1.5, color=colors.HexColor("#1F1C18")),
        Spacer(1, 0.8 * inch),
        Paragraph(f"<b>Submitted by:</b> {team_names or 'Student Name'}", styles["CoverMeta"]),
        Paragraph(f"<b>Guide:</b> {project.get('guide_name', 'Faculty Supervisor')}", styles["CoverMeta"]),
        Paragraph(f"<b>Academic Year:</b> {project.get('academic_year', '2024-25')}", styles["CoverMeta"]),
        PageBreak(),
    ]
    return elements


def _pdf_chapter(chap: dict, styles) -> list:
    elements: list = []
    chapter_num = chap['chapter_number']
    chapter_title_text = chap['chapter_title']
    main_title = f"Chapter {chapter_num}: {chapter_title_text}"

    # Main Heading: 16 pt (Bold)
    elements.append(Paragraph(_markdown_to_pdf_xml(main_title), styles["MainHeading"]))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CCCCCC")))
    elements.append(Spacer(1, 0.2 * inch))

    raw_content = chap.get("content", "")
    blocks = _parse_markdown_blocks(raw_content, strip_duplicate_h1=main_title)

    for kind, payload in blocks:
        if kind == "heading2":
            elements.append(Paragraph(_markdown_to_pdf_xml(payload), styles["SubHeading"]))
        elif kind == "heading3":
            elements.append(Paragraph(_markdown_to_pdf_xml(payload), styles["SubHeading"]))
        elif kind == "bullet":
            elements.append(Paragraph(f"&bull; {_markdown_to_pdf_xml(payload)}", styles["BulletJustified"]))
        elif kind == "code":
            clean_code = _sanitize_unicode_for_pdf(payload)
            escaped_code = clean_code.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            escaped_code = escaped_code.replace("\n", "<br/>").replace(" ", "&nbsp;")
            elements.append(Paragraph(escaped_code, styles["CodeBlock"]))
        elif kind == "table":
            headers = payload.get("headers", [])
            rows = payload.get("rows", [])
            if headers or rows:
                table_data = []
                if headers:
                    table_data.append([
                        Paragraph(_markdown_to_pdf_xml(h), styles["TableCellBold"]) for h in headers
                    ])
                for r in rows:
                    table_data.append([
                        Paragraph(_markdown_to_pdf_xml(c), styles["TableCell"]) for c in r
                    ])
                
                # Full printable width A4 = 487 pt
                num_cols = max(len(headers), max((len(r) for r in rows), default=1))
                col_w = 487.0 / num_cols if num_cols > 0 else 487.0
                col_widths = [col_w] * num_cols

                # Render ReportLab Table
                t = Table(table_data, colWidths=col_widths)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F2ECE1')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#1F1C18')),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D3C9B8')),
                ]))
                elements.append(Spacer(1, 0.1 * inch))
                elements.append(t)
                elements.append(Spacer(1, 0.1 * inch))
        else:
            # Body Paragraph: 12 pt (Justified)
            elements.append(Paragraph(_markdown_to_pdf_xml(payload), styles["BodyJustified"]))

    elements.append(PageBreak())
    return elements


def _sanitize_unicode_for_pdf(text: str) -> str:
    """Replaces Unicode characters not supported by ReportLab standard fonts with ASCII equivalents to eliminate black square boxes (■)."""
    if not text:
        return ""

    # 1. Replace Box Drawing & Geometric Symbols (e.g. └─ ├─ │ ─ █ ■ ▸ →)
    symbol_map = {
        # Boxes & Geometric Shapes
        '█': '#', '■': '*', '□': '*', '▪': '*', '▫': '*', '▬': '-', '░': ' ', '▒': ' ', '▓': '#',
        '▲': '^', '▼': 'v', '►': '>', '◄': '<', '▸': '>', '◂': '<', '▶': '>', '◀': '<',
        '●': '*', '○': 'o', '◆': '*', '◇': 'o', '◼': '*', '◾': '*', '▮': '|', '▯': '|', '▰': '-', '▱': '-',
        '◈': '*', '◉': '*', '◊': 'o', '◌': 'o', '◍': 'o', '◎': 'o', '◐': 'o', '◑': 'o',

        # Box Drawing & Tree Diagrams
        '│': '|', '┃': '|', '║': '|', '┆': '|', '┇': '|', '┊': '|', '┈': '-', '┉': '-',
        '─': '-', '━': '-', '═': '-', '┄': '-', '┅': '-',
        '┌': '+', '┍': '+', '┎': '+', '┏': '+', '╒': '+', '╓': '+', '╔': '+',
        '┐': '+', '┑': '+', '┒': '+', '┓': '+', '╕': '+', '╖': '+', '╗': '+',
        '└': '+', '┕': '+', '┖': '+', '┗': '+', '╘': '+', '╙': '+', '╚': '+',
        '┘': '+', '┙': '+', '┚': '+', '┛': '+', '╛': '+', '╜': '+', '╝': '+',
        '├': '+', '┝': '+', '┞': '+', '┟': '+', '┠': '+', '┡': '+', '┢': '+', '┣': '+', '╞': '+', '╟': '+', '╠': '+',
        '┤': '+', '┥': '+', '┦': '+', '┧': '+', '┨': '+', '┩': '+', '┪': '+', '┫': '+', '╡': '+', '╢': '+', '╣': '+',
        '┬': '+', '┭': '+', '┮': '+', '┯': '+', '┰': '+', '┱': '+', '┲': '+', '┳': '+', '╤': '+', '╥': '+', '╦': '+',
        '┴': '+', '┵': '+', '┶': '+', '┷': '+', '┸': '+', '┹': '+', '┺': '+', '┻': '+', '╧': '+', '╨': '+', '╩': '+',
        '┼': '+', '┽': '+', '┾': '+', '┿': '+', '╀': '+', '╁': '+', '╂': '+', '╃': '+', '╄': '+', '╅': '+', '╆': '+', '╇': '+', '╈': '+', '╉': '+', '╊': '+', '╋': '+', '╪': '+', '╫': '+', '╬': '+',

        # Dashes & Hyphens
        '\u2010': '-', '\u2011': '-', '\u2012': '-', '\u2013': '-', '\u2014': '--', '\u2015': '--', '\u2212': '-',

        # Arrows
        '→': '->', '➔': '->', '➜': '->', '➡': '->', '➢': '->', '➣': '->', '➤': '->',
        '➥': '->', '➦': '->', '➪': '->', '➩': '->', '➬': '->', '➮': '->', '➯': '->',
        '➱': '->', '➲': '->', '➳': '->', '➺': '->', '➼': '->', '➽': '->', '⇒': '=>', '⇔': '<=>',
        '←': '<-', '⬅': '<-', '⇐': '<-', '↔': '<->', '↑': '^', '↓': 'v',

        # Quotes, Bullets, Checkmarks
        '\u2018': "'", '\u2019': "'", '\u201a': "'", '\u201b': "'",
        '\u201c': '"', '\u201d': '"', '\u201e': '"', '\u201f': '"',
        '\u2022': '*', '\u2023': '*', '\u2043': '*', '\u2026': '...',
        '\u00a0': ' ',
        '✔': '[X]', '✓': '[X]', '✅': '[X]', '✖': '[ ]', '✗': '[ ]', '✘': '[ ]', '❌': '[ ]',
    }
    for u_char, ascii_char in symbol_map.items():
        text = text.replace(u_char, ascii_char)

    # 2. Normalize remaining non-ASCII Unicode characters safely to ASCII
    cleaned_chars = []
    for ch in text:
        ord_ch = ord(ch)
        if ord_ch < 128 or ord_ch in (160, 161, 162, 163, 169, 174, 176, 177, 181, 187, 191):
            cleaned_chars.append(ch)
        else:
            norm = unicodedata.normalize('NFKD', ch)
            ascii_chunk = ''.join(c for c in norm if ord(c) < 128)
            if ascii_chunk:
                cleaned_chars.append(ascii_chunk)
            else:
                cleaned_chars.append(" ")

    return "".join(cleaned_chars)


def _markdown_to_pdf_xml(text: str) -> str:
    """Converts markdown inline syntax to ReportLab XML tags safely."""
    if not text:
        return ""
    
    # Sanitize unicode hyphens/quotes/symbols that cause black square boxes (■) in ReportLab
    text = _sanitize_unicode_for_pdf(text)

    # Escape XML special chars
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = text.replace("<br>", "<br/>")

    # Protect inline code `code` from italic/bold underscore replacements using safe token without underscores
    code_tokens = []
    def save_code(match):
        code_tokens.append(match.group(1))
        return f"XCODEPH{len(code_tokens)-1}X"

    text = re.sub(r"`(.*?)`", save_code, text)

    # Convert bold **text** or __text__ -> <b>text</b>
    text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"__(.*?)__", r"<b>\1</b>", text)

    # Convert italic *text* -> <i>text</i>
    text = re.sub(r"(?<!\w)\*(.*?)\*(?!\w)", r"<i>\1</i>", text)

    # Restore code placeholders as Courier font
    for idx, code_str in enumerate(code_tokens):
        text = text.replace(f"XCODEPH{idx}X", f'<font name="Courier">{code_str}</font>')

    return text


# ─── DOCX Helpers ─────────────────────────────────────────────────────────────

def _set_docx_styles(doc: Document) -> None:
    """Apply global DOCX formatting: Times New Roman, 1-inch margins."""
    section = doc.sections[0]
    section.page_height = Inches(11.69)
    section.page_width = Inches(8.27)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x1C, 0x18)


def _docx_cover_page(doc: Document, project: dict) -> None:
    for _ in range(3):
        doc.add_paragraph()

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run(project.get("institution", "University Name"))
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(14)
    r_inst.bold = True

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dept = p_dept.add_run(project.get("department", "Department of Computer Science"))
    r_dept.font.name = 'Times New Roman'
    r_dept.font.size = Pt(14)

    doc.add_paragraph()
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(project.get("title", "Project Title"))
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.bold = True

    doc.add_paragraph()
    team = "; ".join(m.get("name", "") for m in (project.get("team_members") or []))
    
    _add_formatted_docx_paragraph(doc, f"Submitted by: {team or 'Student Name'}", size_pt=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_formatted_docx_paragraph(doc, f"Guide: {project.get('guide_name', 'Faculty Supervisor')}", size_pt=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_formatted_docx_paragraph(doc, f"Academic Year: {project.get('academic_year', '2024-25')}", size_pt=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    doc.add_page_break()


def _docx_chapter(doc: Document, chap: dict) -> None:
    chapter_num = chap['chapter_number']
    chapter_title_text = chap['chapter_title']
    main_title = f"Chapter {chapter_num}: {chapter_title_text}"

    # Main Heading: 16 pt (Bold)
    p_main = doc.add_paragraph()
    p_main.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_main = p_main.add_run(main_title)
    r_main.font.name = 'Times New Roman'
    r_main.font.size = Pt(16)
    r_main.bold = True
    p_main.paragraph_format.space_after = Pt(12)
    p_main.paragraph_format.space_before = Pt(18)

    raw_content = chap.get("content", "")
    blocks = _parse_markdown_blocks(raw_content, strip_duplicate_h1=main_title)

    for kind, payload in blocks:
        if kind in ("heading2", "heading3"):
            # Sub Headings: 14 pt (Bold)
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_sub = p_sub.add_run(payload)
            r_sub.font.name = 'Times New Roman'
            r_sub.font.size = Pt(14)
            r_sub.bold = True
            p_sub.paragraph_format.space_before = Pt(14)
            p_sub.paragraph_format.space_after = Pt(6)

        elif kind == "bullet":
            _add_formatted_docx_paragraph(doc, f"• {payload}", size_pt=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

        elif kind == "code":
            p_code = doc.add_paragraph()
            r_code = p_code.add_run(payload)
            r_code.font.name = 'Courier New'
            r_code.font.size = Pt(11)
            p_code.paragraph_format.space_after = Pt(8)

        elif kind == "table":
            headers = payload.get("headers", [])
            rows = payload.get("rows", [])
            if headers or rows:
                num_rows = len(rows) + (1 if headers else 0)
                num_cols = max(len(headers), max((len(r) for r in rows), default=0))
                if num_rows > 0 and num_cols > 0:
                    tbl = doc.add_table(rows=num_rows, cols=num_cols)
                    tbl.style = 'Table Grid'
                    
                    row_offset = 0
                    if headers:
                        for c_idx, h_text in enumerate(headers):
                            cell = tbl.cell(0, c_idx)
                            cell.text = ""
                            p_cell = cell.paragraphs[0]
                            r_cell = p_cell.add_run(h_text)
                            r_cell.font.name = 'Times New Roman'
                            r_cell.font.size = Pt(12)
                            r_cell.bold = True
                        row_offset = 1

                    for r_idx, r_data in enumerate(rows):
                        for c_idx, val_text in enumerate(r_data):
                            if c_idx < num_cols:
                                cell = tbl.cell(r_idx + row_offset, c_idx)
                                cell.text = ""
                                p_cell = cell.paragraphs[0]
                                r_cell = p_cell.add_run(val_text)
                                r_cell.font.name = 'Times New Roman'
                                r_cell.font.size = Pt(12)

        else:
            # Content Paragraph: 12 pt (Justified)
            _add_formatted_docx_paragraph(doc, payload, size_pt=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    doc.add_page_break()


def _add_formatted_docx_paragraph(doc: Document, text: str, size_pt: int = 12, bold_all: bool = False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Adds a paragraph parsing inline markdown (**bold**, *italic*, `code`) with Times New Roman font."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15

    # Tokenize inline markdown
    tokens = re.split(r"(\*\*.*?\*\*|__.*?__|\*.*?\*|_.*?_|`.*?`)", text)
    for token in tokens:
        if not token:
            continue
        if (token.startswith("**") and token.endswith("**")) or (token.startswith("__") and token.endswith("__")):
            r = p.add_run(token[2:-2])
            r.bold = True
        elif (token.startswith("*") and token.endswith("*")) or (token.startswith("_") and token.endswith("_")):
            r = p.add_run(token[1:-1])
            r.italic = True
        elif token.startswith("`") and token.endswith("`"):
            r = p.add_run(token[1:-1])
            r.font.name = 'Courier New'
        else:
            r = p.add_run(token)

        r.font.name = 'Times New Roman'
        r.font.size = Pt(size_pt)
        if bold_all:
            r.bold = True

    return p


# ─── Markdown Parser Core ──────────────────────────────────────────────────────

def _parse_markdown_blocks(content: str, strip_duplicate_h1: str | None = None) -> list[tuple[str, any]]:
    """
    Parses Markdown into structured blocks:
      - ("heading2", text)
      - ("heading3", text)
      - ("paragraph", text)
      - ("bullet", text)
      - ("table", {"headers": [...], "rows": [[...], ...]})
      - ("code", text)
    """
    lines = content.splitlines()
    blocks = []
    
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i].rstrip()
        stripped = line.strip()

        # Skip empty lines
        if not stripped:
            i += 1
            continue

        # Skip horizontal dividers
        if stripped in ("---", "***", "___"):
            i += 1
            continue

        # Check for Code Blocks ``` ... ```
        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < n:
                i += 1  # consume closing ```
            blocks.append(("code", "\n".join(code_lines)))
            continue

        # Check for Markdown Tables | col1 | col2 |
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table lines into headers and rows
            parsed_table = _parse_markdown_table_lines(table_lines)
            if parsed_table:
                blocks.append(("table", parsed_table))
            continue

        # Check Headings
        if stripped.startswith("# "):
            title_text = stripped[2:].strip()
            # If this heading matches the main chapter title, skip it to prevent duplication
            if strip_duplicate_h1 and (strip_duplicate_h1.lower() in title_text.lower() or title_text.lower() in strip_duplicate_h1.lower()):
                i += 1
                continue
            blocks.append(("heading2", title_text))
            i += 1
            continue
        elif stripped.startswith("## "):
            blocks.append(("heading2", stripped[3:].strip()))
            i += 1
            continue
        elif stripped.startswith("### "):
            blocks.append(("heading3", stripped[4:].strip()))
            i += 1
            continue

        # Check Bullets / Lists
        if stripped.startswith(("- ", "* ", "• ")) or re.match(r"^\d+\.\s+", stripped):
            item_text = re.sub(r"^(-\s*|\*\s*|•\s*|\d+\.\s*)", "", stripped)
            blocks.append(("bullet", item_text))
            i += 1
            continue

        # Regular Paragraph
        blocks.append(("paragraph", stripped))
        i += 1

    return blocks


def _parse_markdown_table_lines(lines: list[str]) -> dict | None:
    """Parses markdown table lines into headers and data rows."""
    if not lines:
        return None
    
    rows_data = []
    for line in lines:
        # Split by | and filter empty leading/trailing cells
        cells = [c.strip() for c in line.split("|")]
        if cells and cells[0] == "":
            cells = cells[1:]
        if cells and cells[-1] == "":
            cells = cells[:-1]
        rows_data.append(cells)

    if not rows_data:
        return None

    headers = rows_data[0]
    data_rows = []

    for r in rows_data[1:]:
        # Skip separator row (e.g., |---|---|)
        if all(re.match(r"^:?-+:?$", c) for c in r if c):
            continue
        data_rows.append(r)

    return {"headers": headers, "rows": data_rows}


def _filter_chapters(chapters: list[dict], include: list[int] | None) -> list[dict]:
    if include is None:
        return sorted(chapters, key=lambda c: c["chapter_number"])
    return sorted(
        [c for c in chapters if c["chapter_number"] in include],
        key=lambda c: c["chapter_number"],
    )


def _safe_filename(title: str) -> str:
    safe = re.sub(r"[^\w\s-]", "", title)
    safe = re.sub(r"\s+", "_", safe).strip("-_")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{safe}_{timestamp}"
