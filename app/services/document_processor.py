"""
Document processing service.

Handles:
  - Text extraction from PDF, DOCX, PPTX, source-code files, and images
  - Text cleaning and normalisation
  - Chunking for embedding
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Iterator

import chardet
import pdfplumber
from docx import Document as DocxDocument
from loguru import logger
from PIL import Image
from pptx import Presentation


# ─── Public API ───────────────────────────────────────────────────────────────

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".doc",
    ".pptx", ".ppt",
    ".txt", ".py", ".js", ".ts", ".java", ".c", ".cpp", ".cs",
    ".html", ".css", ".json", ".xml", ".yaml", ".yml",
    ".md", ".rst",
    ".png", ".jpg", ".jpeg", ".bmp", ".tiff",
}


def extract_text(file_path: str | Path, file_bytes: bytes | None = None) -> str:
    """
    Extract plain text from a file.

    Parameters
    ----------
    file_path : path to the file (used to determine the file type)
    file_bytes : raw bytes of the file content (if already in memory)

    Returns
    -------
    Extracted text as a single string.
    """
    path = Path(file_path)
    ext = path.suffix.lower()
    content = file_bytes or path.read_bytes()

    try:
        if ext == ".pdf":
            return _extract_pdf(content)
        elif ext in (".docx", ".doc"):
            return _extract_docx(content)
        elif ext in (".pptx", ".ppt"):
            return _extract_pptx(content)
        elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".tiff"):
            return _extract_image(content)
        elif ext in SUPPORTED_EXTENSIONS:
            return _extract_text_file(content)
        else:
            logger.warning(f"Unsupported file type: {ext}")
            return ""
    except Exception as exc:
        logger.error(f"Text extraction failed for {path.name}: {exc}")
        return ""


def clean_text(text: str) -> str:
    """Remove noise from extracted text."""
    # Collapse multiple blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Remove non-printable characters (keep newlines/tabs)
    text = re.sub(r"[^\x09\x0A\x0D\x20-\x7E\u00A0-\uFFFF]", "", text)
    # Strip leading/trailing whitespace per line
    lines = [ln.strip() for ln in text.splitlines()]
    text = "\n".join(lines)
    return text.strip()


def chunk_text(
    text: str,
    chunk_size: int = 800,
    chunk_overlap: int = 150,
) -> list[str]:
    """
    Split text into overlapping chunks suitable for embedding.

    Uses a simple sentence-aware splitter that tries to break on paragraph
    boundaries first, then sentence boundaries.
    """
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return splitter.split_text(text)


def process_file(file_path: str | Path, file_bytes: bytes | None = None) -> list[str]:
    """
    Full pipeline: extract → clean → chunk.
    Returns a list of text chunks ready for embedding.
    """
    raw = extract_text(file_path, file_bytes)
    cleaned = clean_text(raw)
    if not cleaned:
        return []
    chunks = chunk_text(cleaned)
    logger.info(f"Processed '{Path(file_path).name}' → {len(chunks)} chunks")
    return chunks


# ─── Private Helpers ──────────────────────────────────────────────────────────

def _extract_pdf(content: bytes) -> str:
    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                parts.append(page_text)
    return "\n\n".join(parts)


def _extract_docx(content: bytes) -> str:
    doc = DocxDocument(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _extract_pptx(content: bytes) -> str:
    prs = Presentation(io.BytesIO(content))
    slides_text: list[str] = []
    for slide in prs.slides:
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                slide_parts.append(shape.text_frame.text)
        if slide_parts:
            slides_text.append(" ".join(slide_parts))
    return "\n\n".join(slides_text)


def _extract_image(content: bytes) -> str:
    """OCR-based extraction using pytesseract (optional — returns empty if not installed)."""
    try:
        import pytesseract
        img = Image.open(io.BytesIO(content))
        return pytesseract.image_to_string(img)
    except ImportError:
        logger.warning("pytesseract not installed – skipping image OCR")
        return ""
    except Exception as exc:
        logger.warning(f"Image OCR failed: {exc}")
        return ""


def _extract_text_file(content: bytes) -> str:
    detected = chardet.detect(content)
    encoding = detected.get("encoding") or "utf-8"
    try:
        return content.decode(encoding, errors="replace")
    except Exception:
        return content.decode("utf-8", errors="replace")
